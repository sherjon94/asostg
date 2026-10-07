# -*- coding: utf-8 -*-
"""Parse Antiplagiat (antiplag.uz) PDF reports and build an Asosnoma .docx.

Supports output in 4 variants: uz-cyrl, uz-latn, ru, en.
"""
import io
import re
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from . import translit

# ---------------------------------------------------------------- categories
EXCLUDED, CITATION, BIBLIO, COMMON = 'excluded', 'citation', 'biblio', 'common'

SCI_DOM = [
    'springer', 'pmc.ncbi', 'frontiersin', 'ieee', 'mdpi', 'biomedcentral',
    'kjronline', 'sciencedirect', 'nature', 'science-medicine', 'uniroma',
    'haik.upjs', 'elibrary', 'cyberleninka', 'ncbi.nlm', 'docplayer',
    'radiology24', 'collections.nlm', 'link.springer'
]
DISS_DOM = [
    'dlib.rsl.ru', 'diss.natlib.uz', 'vuzring', 'inlibrary.uz', 'ziyonet.uz',
    'repository.tma', 'natmedlib', 'fips.ru', 'ir.lib.uwo', 'repub.eur',
    'cancercenter', 'sammu', 'gpmu', 'adti.uz', 'jizpi'
]

# -------------------------------------------------------------- localisation
L10N = {
    'uz-cyrl': {
        'title': 'АСОСНОМА',
        'col_no': '№', 'col_src': 'Электрон матн кўринишида', 'col_note': 'Изоҳ',
        'sec': {
            'uz': 'Ўзбек тилидаги матнидан',
            'ru': 'Рус тилидаги матнидан',
            'eng': 'Инглиз тилидаги матнидан',
            'ai': 'Сунъий интеллект текширувидаги матнидан'
        },
        EXCLUDED: 'Манба истисно қилинган. Сабаби: Кесишманинг кичик фоизи.',
        CITATION: 'Иқтибослик келтирилган.',
        BIBLIO: 'Мазкур матн фойдаланилган адабиётлар рўйхатида берилган.',
        COMMON: 'Умум қабул қилинган сўзлар бирикмаси келтирилган.',
        'header': ('{uni} {status} {fio}нинг «{topic}» мавзусидаги {spec} '
                   'ихтисосликлари бўйича тайёрланган {degree} юзасидан плагиатга '
                   'текширувчи дастурда мос деб топилган камчиликларга'),
        'ai_note': ('СИ – {ai}%; эҳтимолий СИ – {prob}%; оригинал матн – {orig}%. Тизим '
                    'томонидан белгиланган матн қисмлари илмий-тиббий атамалар ҳамда умум '
                    'қабул қилинган ибора ва таърифлардан иборат бўлиб, диссертация матни '
                    'муаллиф томонидан мустақил ёзилган. Сунъий интеллект детектори ёрдамчи '
                    'восита бўлиб, унинг кўрсаткичлари эҳтимолий хусусиятга эга.'),
        'ai_src': 'Сунъий интеллект детектори натижаси',
        'comm_head': 'ЭКСПЕРТ КОМИССИЯСИ АЪЗОЛАРИ:',
    },
    'uz-latn': {
        'title': 'ASOSNOMA',
        'col_no': '№', 'col_src': 'Elektron matn koʻrinishida', 'col_note': 'Izoh',
        'sec': {
            'uz': 'Oʻzbek tilidagi matnidan',
            'ru': 'Rus tilidagi matnidan',
            'eng': 'Ingliz tilidagi matnidan',
            'ai': 'Sunʼiy intellekt tekshiruvidagi matnidan'
        },
        EXCLUDED: 'Manba istisno qilingan. Sababi: Kesishmaning kichik foizi.',
        CITATION: 'Iqtiboslik keltirilgan.',
        BIBLIO: 'Mazkur matn foydalanilgan adabiyotlar roʻyxatida berilgan.',
        COMMON: 'Umum qabul qilingan soʻzlar birikmasi keltirilgan.',
        'header': ('{uni} {status} {fio}ning «{topic}» mavzusidagi {spec} '
                   'ixtisosliklari boʻyicha tayyorlangan {degree} yuzasidan plagiatga '
                   'tekshiruvchi dasturda mos deb topilgan kamchiliklarga'),
        'ai_note': ('SI – {ai}%; ehtimoliy SI – {prob}%; original matn – {orig}%. Tizim '
                    'tomonidan belgilangan matn qismlari ilmiy-tibbiy atamalar hamda umum '
                    'qabul qilingan ibora va taʼriflardan iborat boʻlib, dissertatsiya matni '
                    'muallif tomonidan mustaqil yozilgan. Sunʼiy intellekt detektori yordamchi '
                    'vosita boʻlib, uning koʻrsatkichlari ehtimoliy xususiyatga ega.'),
        'ai_src': 'Sunʼiy intellekt detektori natijasi',
        'comm_head': 'EKSPERT KOMISSIYASI AʼZOLARI:',
    },
    'ru': {
        'title': 'ОБОСНОВАНИЕ',
        'col_no': '№', 'col_src': 'В виде электронного текста', 'col_note': 'Примечание',
        'sec': {
            'uz': 'Из текста на узбекском языке',
            'ru': 'Из текста на русском языке',
            'eng': 'Из текста на английском языке',
            'ai': 'Из текста проверки на искусственный интеллект'
        },
        EXCLUDED: 'Источник исключён. Причина: малый процент совпадения.',
        CITATION: 'Приведено цитирование.',
        BIBLIO: 'Данный текст приведён в списке использованной литературы.',
        COMMON: 'Приведены общепринятые словосочетания.',
        'header': ('По диссертации {status} {uni} {fio} на тему «{topic}» по '
                   'специальностям {spec}, подготовленной на соискание {degree}, — по '
                   'недостаткам, признанным совпадениями программой проверки на плагиат'),
        'ai_note': ('ИИ – {ai}%; вероятный ИИ – {prob}%; оригинальный текст – {orig}%. '
                    'Отмеченные системой фрагменты состоят из научно-медицинских терминов и '
                    'общепринятых выражений; текст диссертации написан автором самостоятельно. '
                    'Детектор ИИ является вспомогательным средством, его показатели носят '
                    'вероятностный характер.'),
        'ai_src': 'Результат детектора искусственного интеллекта',
        'comm_head': 'ЧЛЕНЫ ЭКСПЕРТНОЙ КОМИССИИ:',
    },
    'en': {
        'title': 'JUSTIFICATION (ASOSNOMA)',
        'col_no': 'No.', 'col_src': 'In electronic text form', 'col_note': 'Comment',
        'sec': {
            'uz': 'From the text in Uzbek',
            'ru': 'From the text in Russian',
            'eng': 'From the text in English',
            'ai': 'From the AI-detection text'
        },
        EXCLUDED: 'The source is excluded. Reason: small percentage of overlap.',
        CITATION: 'A citation is provided.',
        BIBLIO: 'This text is included in the list of references.',
        COMMON: 'Commonly accepted word combinations are used.',
        'header': ('This justification concerns the items found as matches by the '
                   'anti-plagiarism checking software in the {degree} of {fio}, {status} '
                   'of {uni}, on the topic “{topic}” in the specialities {spec}'),
        'ai_note': ('AI – {ai}%; probable AI – {prob}%; original text – {orig}%. The '
                    'fragments flagged by the system consist of scientific and medical terms '
                    'and commonly accepted expressions; the dissertation text was written '
                    'independently by the author. The AI detector is an auxiliary tool and its '
                    'indicators are of a probabilistic nature.'),
        'ai_src': 'Artificial intelligence detector result',
        'comm_head': 'MEMBERS OF THE EXPERT COMMISSION:',
    },
}

# status in the header (genitive form) per language
STATUS = {
    'uz-cyrl': {'tayanch': 'таянч докторанти', 'mustaqil': 'мустақил изланувчиси'},
    'uz-latn': {'tayanch': 'tayanch doktoranti', 'mustaqil': 'mustaqil izlanuvchisi'},
    'ru': {'tayanch': 'базового докторанта', 'mustaqil': 'самостоятельного соискателя'},
    'en': {'tayanch': 'basic doctoral student', 'mustaqil': 'independent researcher'},
}

# status as the signature-block label (nominative form)
STATUS_SIGNER = {
    'uz-cyrl': {'tayanch': 'Таянч докторант', 'mustaqil': 'Мустақил изланувчи'},
    'uz-latn': {'tayanch': 'Tayanch doktorant', 'mustaqil': 'Mustaqil izlanuvchi'},
    'ru': {'tayanch': 'Базовый докторант', 'mustaqil': 'Самостоятельный соискатель'},
    'en': {'tayanch': 'Basic doctoral student', 'mustaqil': 'Independent researcher'},
}

DEFAULT_COMMISSION = [
    ['Илмий даражалар берувчи илмий кенгаш раиси', ''],
    ['Илмий даражалар берувчи илмий кенгаш котиби', ''],
    ['Илмий даражалар берувчи илмий кенгаш аъзоси', ''],
    ['Илмий раҳбар', ''],
    ['Илмий раҳбар', ''],
    ['Таянч докторант', ''],
]

DEFAULT_META = {
    'fio': '',
    'topic': '',
    'spec': '',
    'university': '',
    'degree': 'фалсафа доктори (PhD) диссертацияси',
    'status': 'tayanch',
    'commission': DEFAULT_COMMISSION,
}


def format_short_name(full_name: str) -> str:
    """e.g. 'Гайбуллаев Шерзод Обид ўғли' -> 'Ш.О. Гайбуллаев'"""
    if not full_name:
        return ''
    parts = full_name.strip().split()
    if len(parts) >= 3 and parts[-1].lower() in ('ўғли', 'угли', 'кизи', 'қизи', 'o\'g\'li', 'oʻgʻli', 'qizi'):
        # Last is ugli/qizi, parts[0]=Surname, parts[1]=First, parts[2]=Middle
        surname = parts[0]
        i1 = parts[1][0].upper() + '.'
        i2 = parts[2][0].upper() + '.'
        return f"{i1}{i2} {surname}"
    elif len(parts) >= 2:
        surname = parts[0]
        i1 = parts[1][0].upper() + '.'
        i2 = (parts[2][0].upper() + '.') if len(parts) >= 3 else ''
        return f"{i1}{i2} {surname}".strip()
    return full_name


# ------------------------------------------------------------------- parsing
def parse_report(pdf_bytes):
    """Return list of sources: {num, matn, hisob, name, url, excluded}."""
    doc = fitz.open(stream=pdf_bytes, filetype='pdf')
    sources = {}
    for pno in range(doc.page_count):
        pg = doc[pno]
        words = pg.get_text('words')
        # Match anchors like [1], [01], [25]
        anchors = [w for w in words if w[0] < 65 and re.fullmatch(r'\[\d+\]', w[4])]
        if not anchors:
            continue
        anchors.sort(key=lambda w: w[1])

        # Allow both dot and comma in percentage: e.g. 0.55%, 0,55%, 3%, 3.0%
        matn = [w for w in words if 45 <= w[0] <= 115 and re.fullmatch(r'\d+([.,]\d+)?\s*%', w[4])]
        matn.sort(key=lambda w: w[1])
        matn_by = {}
        for a in anchors:
            cands = [m for m in matn if abs(m[1] - a[1]) <= 12]
            num = int(a[4].strip('[]'))
            matn_by[num] = cands[0][4] if cands else ''

        txt = pg.get_text('text')
        marks = list(re.finditer(r'\[(\d+)\]', txt))
        for i, m in enumerate(marks):
            num = int(m.group(1))
            seg = txt[m.end(): marks[i + 1].start() if i + 1 < len(marks) else len(txt)]
            lines = [l.strip() for l in seg.splitlines() if l.strip()]
            hisob = lines[0] if lines and re.fullmatch(r'\d+([.,]\d+)?\s*%', lines[0]) else ''
            rest = lines[1:] if hisob else lines
            urls = [l for l in rest if l.startswith('http')]
            url = urls[0] if urls else ''
            excluded = any('istisno' in l.lower() or 'kesishma' in l.lower() for l in rest)
            name_parts = []
            for l in rest:
                if l.startswith('http'):
                    break
                if re.match(r'\d{1,2}\s', l) and re.search(
                        r'(Yanv|Fev|Mar|Маr|Apr|May|Iyun|Iyul|Avg|Sen|Okt|Noy|Dek)', l):
                    break
                if l in ('unilibrary', 'vuzring', 'media', 'springer', 'ieee', 'elibrary',
                         'garant_documents'):
                    break
                if 'istisno' in l.lower() or 'kesishma' in l.lower() or l == 'Sababi:':
                    break
                name_parts.append(l)
            name = ' '.join(name_parts).strip()
            
            # Prefer matn, fallback to hisob if matn is empty
            final_matn = matn_by.get(num, '') or hisob

            if num not in sources:
                sources[num] = dict(num=num, matn=final_matn, hisob=hisob,
                                    name=name, url=url, excluded=excluded)
    return [sources[k] for k in sorted(sources)]


def parse_ai(pdf_bytes):
    """Return (ai, probable, original) percentages as strings, or None."""
    try:
        doc = fitz.open(stream=pdf_bytes, filetype='pdf')
        if doc.page_count == 0:
            return None
        txt = doc[0].get_text()
        found = re.findall(r'(\d{1,3}(?:[.,]\d{1,2})?)\s*%', txt)
        if len(found) >= 3:
            return found[0], found[1], found[2]
        found_nums = re.findall(r'(\d{1,3}[.,]\d{1,2})', txt)
        if len(found_nums) >= 3:
            return found_nums[0], found_nums[1], found_nums[2]
    except Exception:
        pass
    return None


def extract_meta_from_pdf(pdf_bytes):
    """Robust multi-strategy extraction of dissertation metadata from any report/dissertation PDF.
    Supports:
    - Strategy 1: Header introductory paragraph (Format A: Usarov, Gaybullayev, etc.)
    - Strategy 2: Vertical dissertation title page (Format B: antiplag reports in uz, ru, eng, etc.)
    Returns a dict with any of: fio, university, topic, spec, degree, status."""
    try:
        doc = fitz.open(stream=pdf_bytes, filetype='pdf')
    except Exception:
        return {}

    meta = {}

    # ==============================================================
    # STRATEGY 1: Search for Header Paragraph (Format A)
    # e.g.: "[Uni] [Status] [FIO]ning «[Topic]» mavzusidagi [Spec] ... [Degree] ... yuzasidan"
    # ==============================================================
    for pno in range(min(15, doc.page_count)):
        txt = doc[pno].get_text()
        norm = re.sub(r'\s+', ' ', txt).strip()

        m_stat = re.search(
            r'\b(таянч\s+докторанти|мустақил\s+изланувчиси|базового\s+докторанта|соискателя|tayanch\s+doktoranti|mustaqil\s+izlanuvchisi|докторанти|doktoranti)\b',
            norm, re.IGNORECASE
        )
        m_topic = re.search(r'[«\"“]([^»\"”]{10,250})[»\"”]', norm)

        if m_stat and m_topic and m_topic.start() > m_stat.end():
            # 1. University: text before status
            uni_candidate = norm[:m_stat.start()].strip()
            for marker in ('ВАЗИРЛИГИ', 'VAZIRLIGI', 'МИНИСТЕРСТВО', 'MINISTRY'):
                if marker in uni_candidate.upper():
                    uni_candidate = uni_candidate[uni_candidate.upper().rindex(marker) + len(marker):].strip()
            uni_candidate = re.sub(r'^[^\w]+', '', uni_candidate).strip()
            if uni_candidate:
                meta['university'] = uni_candidate

            # 2. Status
            st_raw = m_stat.group(1).lower()
            meta['status'] = 'mustaqil' if any(w in st_raw for w in ['мустақил', 'mustaqil', 'соискател']) else 'tayanch'

            # 3. FIO: text between status and topic
            fio_candidate = norm[m_stat.end():m_topic.start()].strip()
            fio_candidate = re.sub(r'(?i)\b(нинг|ning|\s+на\s+тему)\s*$', '', fio_candidate).strip()
            fio_candidate = re.sub(r'(?i)(нинг|ning)$', '', fio_candidate).strip()
            fio_candidate = re.sub(r'(?i)\b(угли|ўғли|ўгли)\b', 'ўғли', fio_candidate)
            fio_candidate = re.sub(r'(?i)\b(кизи|қизи)\b', 'қизи', fio_candidate)
            fio_candidate = re.sub(r'(?i)\b(o[\'ʻ’`]?g[\'ʻ’`]?li)\b', "oʻgʻli", fio_candidate)
            fio_candidate = re.sub(r'(?i)\b(qizi)\b', "qizi", fio_candidate)
            if fio_candidate:
                meta['fio'] = fio_candidate

            # 4. Topic
            meta['topic'] = m_topic.group(1).strip()

            # 5. Specialties and Degree: from text after topic
            post_topic = norm[m_topic.end():]
            specs = re.findall(r'(\d\d\.\d\d\.\d\d)\s*[–\-—:]\s*([A-Za-zА-Яа-яЎўҚқҒғҲҳ\s]+?)(?=,\s*\d\d\.\d\d\.\d\d|\s*ихтисослик|\s*мутахассислик|\s*бўйича|\s*тайёрланган|$)', post_topic)
            if specs:
                parts = []
                for c, n in specs:
                    parts.append(f"{c} – {n.strip()}")
                meta['spec'] = ', '.join(parts)

            deg_low = post_topic.lower()
            if 'dsc' in deg_low or 'фан доктори' in deg_low or 'доктор наук' in deg_low:
                meta['degree'] = 'фан доктори (DSc) диссертацияси'
            else:
                meta['degree'] = 'фалсафа доктори (PhD) диссертацияси'

            if meta.get('fio') and meta.get('topic'):
                return meta

    # ==============================================================
    # STRATEGY 2: Vertical Dissertation Title Page (Titul varaqasi)
    # ==============================================================
    best_page = None
    best_score = -1

    for pno in range(min(35, doc.page_count)):
        t = doc[pno].get_text()
        tl = t.lower()
        score = 0
        if re.search(r'\b\d\d\.\d\d\.\d\d\b', t):
            score += 30
        if any(w in tl for w in ['диссертац', 'dissertat']):
            score += 20
        if any(w in tl for w in ['phd', 'dsc', 'доктор', 'doctor']):
            score += 15
        if any(w in tl for w in ['мутахассислик', 'ихтисослик', 'mutaxassislik', 'ixtisoslik', 'специальност', 'specialty', 'cipher', 'шифр']):
            score += 15
        if any(w in tl for w in ['удк', 'udk', 'udc', 'қўл ёзма', 'қўлёзма', 'qo‘lyozma', 'qo\'lyozma', 'на правах рукописи', 'hand written']):
            score += 15
        if any(w in tl for w in ['университет', 'институт', 'universitet', 'institut', 'university', 'institute']):
            score += 10
        if any(w in tl for w in ['илмий раҳбар', 'илмий маслаҳатчи', 'научный руководитель', 'scientific leader']):
            score += 10

        lines_count = len([l for l in t.splitlines() if l.strip()])
        if 8 <= lines_count <= 60 and score > best_score and score >= 40:
            best_score = score
            best_page = t

    if best_page:
        lines = [l.strip() for l in best_page.splitlines() if l.strip()]
        low = best_page.lower()

        # 1. University
        for l in lines:
            up = l.upper()
            if any(w in up for w in ['УНИВЕРСИТЕТ', 'ИНСТИТУТ', 'АКАДЕМИЯ', 'UNIVERSITET', 'INSTITUT', 'AKADEMIYA', 'UNIVERSITY', 'INSTITUTE', 'ACADEMY', 'MARKAZ', 'МАРКАЗ']):
                u = l
                for marker in ('ВАЗИРЛИГИ', 'VAZIRLIGI', 'МИНИСТЕРСТВО', 'MINISTRY'):
                    if marker in u.upper():
                        u = u[u.upper().rindex(marker) + len(marker):].strip()
                meta['university'] = u.capitalize() if u.isupper() else u
                break

        # 2. Status
        if any(w in low for w in ['мустақил', 'mustaqil', 'соискател']):
            meta['status'] = 'mustaqil'
        else:
            meta['status'] = 'tayanch'

        # 3. Degree (Check degree specifically around dissertation purpose)
        deg_found = False
        for l in lines:
            lu = l.lower()
            if any(w in lu for w in ['илмий даражаси', 'на соискание', 'илмий даража', 'ученой степени', 'scientific level', 'scientific degree']):
                if 'dsc' in lu or 'фан доктори' in lu or 'доктор наук' in lu:
                    meta['degree'] = 'фан доктори (DSc) диссертацияси'
                    deg_found = True
                    break
                elif 'phd' in lu or 'фалсафа доктори' in lu or 'доктор философии' in lu or 'doctor of philosophy' in lu:
                    meta['degree'] = 'фалсафа доктори (PhD) диссертацияси'
                    deg_found = True
                    break
        if not deg_found:
            meta['degree'] = 'фалсафа доктори (PhD) диссертацияси'

        # 4. Specialities
        specs = re.findall(r'(\d\d\.\d\d\.\d\d)\s*[–\-—:]\s*([^\n;]+)', best_page)
        if specs:
            seen, parts = set(), []
            for code, name in specs:
                name_clean = re.sub(r'(?i)\b(Specialty code|Код специальности|код|шифри|шифр|ихтисослик.*)\b', '', name).strip().rstrip('.,;')
                if code not in seen:
                    seen.add(code)
                    parts.append(f'{code} – {name_clean}')
            meta['spec'] = ', '.join(parts)

        # 5. Author FIO & Topic
        stop_keywords = [
            'мутахассислик', 'ихтисослик', 'специальност', 'specialty', 'шифр', 'cipher',
            'диссертация', 'диссертац', 'dissertation', 'dissertat',
            'фалсафа доктори', 'фан доктори', 'доктор философии', 'доктор наук',
            'илмий раҳбар', 'илмий маслаҳатчи', 'научный руководитель', 'scientific leader',
            'удк', 'udk', 'udc', 'қўл ёзма', 'на правах', 'hand written', 'rights'
        ]

        start_idx = 0
        for i, l in enumerate(lines):
            lu = l.lower()
            if any(w in lu for w in ['удк', 'udk', 'udc', 'рукопис', 'қўл ёзма', 'қўлёзма', 'rights', 'rights :']):
                start_idx = i + 1
            elif any(w in lu for w in ['университет', 'институт', 'universitet', 'institut', 'university']) and start_idx == 0:
                start_idx = i + 1

        fio_idx = None
        for idx in range(start_idx, len(lines)):
            l = lines[idx]
            lu = l.lower()
            if any(k in lu for k in stop_keywords) or re.search(r'\b\d\d\.\d\d\.\d\d\b', l):
                continue
            words = l.split()
            if 2 <= len(words) <= 5:
                is_name_like = all(w[0].isupper() or w.upper() in ('OʻGʻLI', "O'G'LI", 'QIZI', 'УГЛИ', 'КИЗИ', 'ЎҒЛИ') for w in words if w.isalpha())
                has_patronymic = bool(re.search(
                    r'(?i)(ўғли|угли|қизи|кизи|o[\'ʻ’`]?g[\'ʻ’`]?li|qizi|ович|евич|овна|евна|ovich|evich|ovna|yevna)$',
                    l.strip()
                ))
                has_surname = bool(re.search(
                    r'(?i)\b[А-ЯA-Z][a-zа-я]*(?:ов|ова|ев|ева|ин|ина|ov|ova|yev|yeva|iy|eva)\b',
                    l
                ))
                if is_name_like and (has_patronymic or has_surname or l.isupper()):
                    fio = l.title() if l.isupper() else l
                    fio = re.sub(r'(?i)\b(угли|ўғли|ўгли)\b', 'ўғли', fio)
                    fio = re.sub(r'(?i)\b(кизи|қизи)\b', 'қизи', fio)
                    fio = re.sub(r'(?i)\b(o[\'ʻ’`]?g[\'ʻ’`]?li)\b', "oʻgʻli", fio)
                    fio = re.sub(r'(?i)\b(qizi)\b', "qizi", fio)
                    meta['fio'] = fio
                    fio_idx = idx
                    break

        # 6. Topic from Title Page
        if fio_idx is not None:
            topic_lines = []
            for l in lines[fio_idx + 1:]:
                lu = l.lower()
                if any(k in lu for k in stop_keywords) or re.search(r'\b\d\d\.\d\d\.\d\d\b', l):
                    break
                topic_lines.append(l)
            if topic_lines:
                raw_topic = ' '.join(topic_lines).strip()
                topic_clean = re.sub(r'^[«\"“\s]+|[»\"”\s]+$', '', raw_topic)
                topic_clean = re.sub(r'\s+', ' ', topic_clean)
                meta['topic'] = topic_clean.capitalize() if topic_clean.isupper() else topic_clean

    return meta


def categorise(r):
    if r['excluded']:
        return EXCLUDED
    name = r['name'] or ''
    url = r['url'] or ''
    nl, ul = name.lower(), url.lower()
    if 'клиническ' in nl or 'рекомендац' in nl or 'garant' in ul or 'методическ' in nl:
        return CITATION
    if any(x in ul for x in DISS_DOM):
        return COMMON
    if any(s in name for s in ['диссертац', 'ДИССЕРТАЦ', 'Диссертац', 'Автореферат',
                               'автореферат', 'диссертация', 'каталог', 'Каталог',
                               'Disser', 'disser']):
        return COMMON
    if re.search(r'[А-ЯЁ][а-яё]+,\s+[А-ЯЁ][а-яё]+', name):
        return COMMON
    if any(x in ul for x in SCI_DOM):
        return BIBLIO
    if name and not re.search(r'[А-Яа-яёЁ]', name) and re.search(r'[A-Za-z]{4}', name):
        return BIBLIO
    return COMMON


# -------------------------------------------------------------- docx helpers
def _set_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def _cell(cell, lines, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    if isinstance(lines, str):
        lines = [lines]
    first = True
    for ln in lines:
        if not first:
            p = cell.add_paragraph()
        r = p.add_run(ln)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = 'Times New Roman'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        first = False


def _section_row(table, title):
    row = table.add_row()
    a = row.cells[0].merge(row.cells[1]).merge(row.cells[2])
    _cell(a, title, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_bg(a, 'D9D9D9')


def _src_label(r):
    name = (r['name'] or '').strip()
    url = (r['url'] or '').strip()
    if name and url:
        return [name, url]
    if name:
        return [name]
    if url:
        return [url]
    return ['koʻrsatilmagan']


# ------------------------------------------------------------------- builder
def build_docx(reports, ai_pcts, lang='uz-cyrl', meta=None, phrases_override=None):
    """reports: dict like {'uz': [sources], 'ru': [...], 'eng': [...]} (any subset)."""
    t = dict(L10N[lang])
    if phrases_override:
        t.update(phrases_override)
    m = dict(DEFAULT_META)
    if meta:
        m.update({k: v for k, v in meta.items() if v})

    # Bring author/topic metadata into the target script
    def _fix(val):
        if not isinstance(val, str) or not val:
            return val
        if lang in ('uz-latn', 'en'):
            return translit.cyr2lat(val)
        if lang == 'uz-cyrl' and not re.search(r'[А-Яа-яёЁ]', val):
            return translit.lat2cyr(val)
        return val

    for k in ('fio', 'topic', 'spec', 'university', 'degree'):
        m[k] = _fix(m.get(k, ''))
    
    STANDARD_COMM_TITLES = {
        'ru': {
            'Илмий даражалар берувчи илмий кенгаш раиси': 'Председатель научного совета по присуждению ученых степеней',
            'Илмий даражалар берувчи илмий кенгаш котиби': 'Ученый секретарь научного совета по присуждению ученых степеней',
            'Илмий даражалар берувчи илмий кенгаш аъзоси': 'Член научного совета по присуждению ученых степеней',
            'Илмий раҳбар': 'Научный руководитель',
        },
        'en': {
            'Илмий даражалар берувчи илмий кенгаш раиси': 'Chairman of the scientific council awarding academic degrees',
            'Илмий даражалар берувчи илмий кенгаш котиби': 'Secretary of the scientific council awarding academic degrees',
            'Илмий даражалар берувчи илмий кенгаш аъзоси': 'Member of the scientific council awarding academic degrees',
            'Илмий раҳбар': 'Scientific supervisor',
        }
    }

    raw_comm = [(p + ['', ''])[:2] for p in m.get('commission', DEFAULT_COMMISSION)]
    comm = []
    for title, name in raw_comm:
        if lang in STANDARD_COMM_TITLES and title in STANDARD_COMM_TITLES[lang]:
            t_title = STANDARD_COMM_TITLES[lang][title]
        else:
            t_title = _fix(title)
        comm.append([t_title, _fix(name)])

    # Researcher status (tayanch doktorant / mustaqil izlanuvchi)
    status_key = m.get('status') if m.get('status') in ('tayanch', 'mustaqil') else 'tayanch'
    status_word = STATUS[lang][status_key]
    
    # Auto-fill author short name if last commission signer is blank
    if comm:
        comm[-1][0] = STATUS_SIGNER[lang][status_key]
        if not comm[-1][1] and m.get('fio'):
            comm[-1][1] = _fix(format_short_name(m['fio']))
    
    m['commission'] = comm

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(1.5)

    p = doc.add_paragraph()
    run = p.add_run(t['header'].format(uni=m['university'], fio=m['fio'], topic=m['topic'],
                                       spec=m['spec'], degree=m['degree'], status=status_word))
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    pt = doc.add_paragraph()
    rt = pt.add_run(t['title'])
    rt.font.size = Pt(14)
    rt.font.bold = True
    rt.font.name = 'Times New Roman'
    rt._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pt.paragraph_format.space_before = Pt(6)
    pt.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    _cell(hdr[0], t['col_no'], size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(hdr[1], t['col_src'], size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(hdr[2], t['col_note'], size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in hdr:
        _set_bg(c, 'BFBFBF')

    widths = [Cm(1.1), Cm(15.5), Cm(10.0)]
    try:
        hdr_tr = table.rows[0]._tr.get_or_add_trPr()
        hdr_tr.append(OxmlElement('w:tblHeader'))
    except Exception:
        pass

    counts = {}
    for key in ['uz', 'ru', 'eng']:
        if key not in reports or not reports[key]:
            continue
        _section_row(table, t['sec'][key])
        cc = {EXCLUDED: 0, CITATION: 0, BIBLIO: 0, COMMON: 0}
        for i, r in enumerate(reports[key], 1):
            row = table.add_row()
            _cell(row.cells[0], str(i), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell(row.cells[1], _src_label(r), size=10)
            cat = categorise(r)
            cc[cat] = cc.get(cat, 0) + 1
            pct_prefix = f"{r['matn']}  " if r.get('matn') else ""
            _cell(row.cells[2], f"{pct_prefix}{t[cat]}", size=10)
        counts[key] = cc

    if ai_pcts:
        _section_row(table, t['sec']['ai'])
        row = table.add_row()
        _cell(row.cells[0], '1', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(row.cells[1], t['ai_src'], size=10)
        ai, prob, orig = ai_pcts
        _cell(row.cells[2], t['ai_note'].format(ai=ai, prob=prob, orig=orig), size=10)

    for row in table.rows:
        try:
            r_tr = row._tr.get_or_add_trPr()
            r_tr.append(OxmlElement('w:cantSplit'))
        except Exception:
            pass
        for idx, w in enumerate(widths):
            row.cells[idx].width = w

    doc.add_paragraph()
    hp = doc.add_paragraph()
    hr = hp.add_run(t['comm_head'])
    hr.font.size = Pt(12)
    hr.font.bold = True
    hr.font.name = 'Times New Roman'
    hr._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    doc.add_paragraph()

    st = doc.add_table(rows=len(comm), cols=2)
    st.allow_autofit = False
    for ri, pair in enumerate(comm):
        title, name = (pair + ['', ''])[:2]
        lc = st.rows[ri].cells[0]
        rc = st.rows[ri].cells[1]
        _cell(lc, title, size=12, bold=True)
        _cell(rc, name, size=12, bold=True)
        lc.width = Cm(15.5)
        rc.width = Cm(8.0)
        lc.paragraphs[0].paragraph_format.space_after = Pt(14)
        rc.paragraphs[0].paragraph_format.space_after = Pt(14)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf, counts


def extract_phrases_from_sample(docx_bytes):
    """Pull the 4 justification phrases from an uploaded sample .docx."""
    try:
        d = Document(io.BytesIO(docx_bytes))
    except Exception:
        return {}
    texts = []
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                texts.append(c.text)
    blob = '\n'.join(texts)
    out = {}
    for line in blob.splitlines():
        s = line.strip()
        low = s.lower()
        if not s:
            continue
        if ('истисно' in low or 'исключ' in low or 'exclud' in low) and EXCLUDED not in out:
            out[EXCLUDED] = re.sub(r'^\d+(,\d+)?%\s*', '', s)
        elif ('иқтибос' in low or 'цитир' in low or 'citation' in low) and CITATION not in out:
            out[CITATION] = re.sub(r'^\d+(,\d+)?%\s*', '', s)
        elif ('адабиёт' in low or 'литератур' in low or 'referenc' in low) and BIBLIO not in out:
            out[BIBLIO] = re.sub(r'^\d+(,\d+)?%\s*', '', s)
        elif ('умум қабул' in low or 'общеприня' in low or 'commonly' in low) and COMMON not in out:
            out[COMMON] = re.sub(r'^\d+(,\d+)?%\s*', '', s)
    return out
